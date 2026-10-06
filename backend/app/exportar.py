"""Exportacion de reportes aprobados a PDF y presentacion (paso 9, RF-08).

Ambos formatos se arman SOLO con lo que se guardo en la tabla `narrativas`
(la narrativa que aprobo RRHH y los hechos que recibio la IA): no se
recalculan indicadores. Asi lo que se distribuye es exactamente lo que se
reviso, y los valores ocultos por privacidad nunca aparecen (no estan en
los hechos).

La API decide quien puede exportar y exige que el reporte este aprobado;
aqui solo se arma el archivo. Estructura del reporte: seccion 10.1 del PRD.
"""
import io
import math
import os
import re
import unicodedata
from datetime import datetime
from xml.sax.saxutils import escape
from zoneinfo import ZoneInfo

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import (
    Flowable,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.narrativa import MESES_ES

FORMATOS = {
    "pdf": "application/pdf",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}

ZONA = ZoneInfo(os.getenv("ZONA_HORARIA", "America/Mexico_City"))
PIE = "Confidencial: uso interno de la empresa"

# Formulas de la seccion 10.2 del PRD, para el anexo de metodologia.
FORMULAS = {
    "tiempo_contratacion": "Promedio de días entre la apertura y la contratación de las vacantes cubiertas en el mes.",
    "costo_por_contratacion": "Costo total de los procesos ÷ vacantes cubiertas.",
    "cumplimiento_metas": "Metas logradas ÷ metas asignadas × 100.",
    "cobertura_capacitacion": "Empleados activos inscritos en al menos un programa del mes ÷ empleados activos × 100.",
    "tasa_finalizacion": "Inscripciones completadas ÷ inscripciones × 100.",
    "rotacion_total": "Bajas del mes ÷ headcount promedio × 100.",
    "rotacion_voluntaria": "Bajas voluntarias ÷ headcount promedio × 100.",
    "rotacion_involuntaria": "Bajas involuntarias ÷ headcount promedio × 100.",
    "enps": "% de promotores (9–10) − % de detractores (0–6).",
    "indice_productividad": "Horas efectivas ÷ horas disponibles × 100.",
}

# Semaforo: siempre con texto, el color solo acompana (igual que el tablero).
# (etiqueta, color del indicador, fondo)
ESTADOS = {
    "rojo": ("Crítico", "#d03b3b", "#f9e3e3"),
    "amarillo": ("Atención", "#c98a00", "#fdf1d6"),
    "verde": ("En rango", "#0c8a0c", "#e6f4e6"),
    "por_vigilar": ("Por vigilar", "#5b6b80", "#e8f0fb"),
}
CONFIANZA = {"alta": "Confianza alta", "media": "Confianza media", "baja": "Confianza baja"}

TINTA = "#1d2733"
SECUNDARIO = "#5b6573"
BORDE = "#d5dbe3"
ACENTO = "#2b5797"
APROBADO = "#0c8a0c"


# ------------------------------------------------------------------ comunes
def nombre_mes(periodo: str) -> str:
    anio, mes = (int(p) for p in periodo.split("-"))
    return f"{MESES_ES[mes - 1]} {anio}"


def fecha_hora(valor) -> str:
    """'4 oct 2026, 00:22' en la zona horaria de la empresa."""
    if valor is None:
        return "—"
    if isinstance(valor, str):
        valor = datetime.fromisoformat(valor.replace("Z", "+00:00"))
    local = valor.astimezone(ZONA)
    return f"{local.day} {MESES_ES[local.month - 1][:3]} {local.year}, {local:%H:%M}"


def nombre_archivo(trabajo: dict, formato: str) -> str:
    """reporte-rrhh-corporativo-2026-08.pdf (solo ASCII, seguro en cualquier sistema)."""
    area = unicodedata.normalize("NFKD", trabajo["narrativa"]["area"]).encode("ascii", "ignore").decode()
    area = re.sub(r"[^a-z0-9]+", "-", area.lower()).strip("-") or f"area-{trabajo['area_id']}"
    return f"reporte-rrhh-{area}-{trabajo['periodo']}.{formato}"


def estado_visible(h: dict) -> str:
    """Un indicador en verde con senales de deterioro se muestra como 'por vigilar'."""
    return "por_vigilar" if h.get("en_vigilancia") and h["estado"] == "verde" else h["estado"]


def alertas(n: dict) -> list[dict]:
    return [h for h in n["hechos"] if h["estado"] in ("rojo", "amarillo") or h.get("en_vigilancia")]


def variacion(texto: str | None, evolucion: str | None) -> str:
    if not texto:
        return "—"
    return f"{texto} ({evolucion})" if evolucion else texto


def respaldo(ids: list[str], hechos: dict, fuentes: list[str]) -> str:
    citas = "; ".join(f"{hechos[i]['nombre']}: {hechos[i]['valor_texto']}" if i in hechos else i for i in ids)
    return f"Respaldo: {citas}. Fuente: {', '.join(fuentes)}."


def linea_aprobacion(t: dict) -> str:
    return f"Aprobado por {t['revisada_por']} el {fecha_hora(t['revisada_en'])}"


def linea_origen(n: dict) -> str:
    donde = "local" if n["proveedor"] == "ollama" else "nube, datos sintéticos"
    return (
        f"Redactado por IA ({n.get('modelo') or n['proveedor']}, {donde}, prompt {n['version_prompt']}) "
        f"el {fecha_hora(n['generado_en'])}, solo a partir de indicadores calculados por el sistema."
    )


def conteos(n: dict) -> list[tuple[str, int]]:
    c = n["conteo_estados"]
    filas = [("rojo", c["rojo"]), ("amarillo", c["amarillo"]), ("verde", c["verde"])]
    if c.get("por_vigilar"):
        filas.append(("por_vigilar", c["por_vigilar"]))
    return filas


def generar(trabajo: dict, formato: str) -> bytes:
    """Archivo del reporte aprobado: 'pdf' o 'pptx'."""
    if formato == "pdf":
        return pdf(trabajo)
    if formato == "pptx":
        return presentacion(trabajo)
    raise ValueError(f"Formato desconocido: {formato}")


# ---------------------------------------------------------------------- PDF
# Vera viene dentro de ReportLab: se incrusta en el PDF, cubre el espanol y la
# puntuacion tipografica, y se ve igual en Windows, Linux o el hosting.
pdfmetrics.registerFont(TTFont("Vera", "Vera.ttf"))
pdfmetrics.registerFont(TTFont("Vera-Bold", "VeraBd.ttf"))
pdfmetrics.registerFont(TTFont("Vera-Italic", "VeraIt.ttf"))
pdfmetrics.registerFont(TTFont("Vera-BoldItalic", "VeraBI.ttf"))
pdfmetrics.registerFontFamily("Vera", normal="Vera", bold="Vera-Bold", italic="Vera-Italic", boldItalic="Vera-BoldItalic")


def _estilo(nombre, **kw):
    base = {"fontName": "Vera", "fontSize": 9.5, "leading": 13, "textColor": colors.HexColor(TINTA)}
    return ParagraphStyle(nombre, **{**base, **kw})


E = {
    "alcance": _estilo("alcance", fontName="Vera-Bold", fontSize=9, textColor=colors.HexColor(ACENTO)),
    "titulo": _estilo("titulo", fontName="Vera-Bold", fontSize=19, leading=24, spaceBefore=2, spaceAfter=4),
    "h2": _estilo("h2", fontName="Vera-Bold", fontSize=12.5, leading=16, spaceBefore=12, spaceAfter=5,
                  textColor=colors.HexColor(ACENTO), keepWithNext=1),
    "h3": _estilo("h3", fontName="Vera-Bold", fontSize=10, leading=13),
    "texto": _estilo("texto"),
    "resumen": _estilo("resumen", fontSize=10.5, leading=15),
    "cifra": _estilo("cifra", fontName="Vera-Bold", fontSize=20, leading=24),
    "nota": _estilo("nota", fontSize=8, leading=10.5, textColor=colors.HexColor(SECUNDARIO)),
    "nota_der": _estilo("nota_der", fontSize=8, leading=10.5, textColor=colors.HexColor(SECUNDARIO), alignment=TA_RIGHT),
    "celda": _estilo("celda", fontSize=8.5, leading=11),
    "celda_num": _estilo("celda_num", fontSize=8.5, leading=11, alignment=TA_RIGHT),
    "encabezado": _estilo("encabezado", fontName="Vera-Bold", fontSize=8, leading=10,
                          textColor=colors.HexColor(SECUNDARIO)),
}


def _p(texto: str, estilo="texto") -> Paragraph:
    """Parrafo con el texto ESCAPADO: la IA puede escribir '<' o '&'."""
    return Paragraph(escape(str(texto)), E[estilo])


class Insignia(Flowable):
    """Semaforo: circulo de color + etiqueta (nunca solo color)."""

    def __init__(self, estado: str, tamano: float = 8):
        super().__init__()
        self.etiqueta, color, _ = ESTADOS.get(estado, (estado, SECUNDARIO, "#ffffff"))
        self.color = colors.HexColor(color)
        self.tamano = tamano
        self.width = pdfmetrics.stringWidth(self.etiqueta, "Vera-Bold", tamano) + tamano * 1.6
        self.height = tamano * 1.3

    def wrap(self, *_):
        return self.width, self.height

    def draw(self):
        c = self.canv
        r = self.tamano * 0.38
        c.setFillColor(self.color)
        c.circle(r, self.height / 2, r, stroke=0, fill=1)
        c.setFillColor(colors.HexColor(TINTA))
        c.setFont("Vera-Bold", self.tamano)
        c.drawString(r * 2 + self.tamano * 0.4, self.height / 2 - self.tamano * 0.35, self.etiqueta)


class CanvasNumerado(rl_canvas.Canvas):
    """Canvas que conoce el total de paginas al final ('Página 1 de 3')."""

    def __init__(self, *args, pie_izq="", pie_der="", **kwargs):
        super().__init__(*args, **kwargs)
        self._paginas = []
        self._pie = (pie_izq, pie_der)

    def showPage(self):
        self._paginas.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._paginas)
        for estado in self._paginas:
            self.__dict__.update(estado)
            self._dibujar_pie(total)
            super().showPage()
        super().save()

    def _dibujar_pie(self, total):
        ancho, _ = LETTER
        self.setStrokeColor(colors.HexColor(BORDE))
        self.setLineWidth(0.5)
        self.line(2 * cm, 1.45 * cm, ancho - 2 * cm, 1.45 * cm)
        self.setFont("Vera", 7.5)
        self.setFillColor(colors.HexColor(SECUNDARIO))
        self.drawString(2 * cm, 1.0 * cm, self._pie[0])
        self.drawRightString(ancho - 2 * cm, 1.0 * cm, f"{self._pie[1]} · Página {self._pageNumber} de {total}")


def _tabla(filas, anchos, encabezado=True):
    t = Table(filas, colWidths=anchos, repeatRows=1 if encabezado else 0)
    t.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor(BORDE)),
                *([("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.HexColor(TINTA))] if encabezado else []),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return t


def _recuadro(contenido: list, fondo: str, borde: str, ancho: float) -> Table:
    t = Table([[contenido]], colWidths=[ancho])
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(fondo)),
                ("LINEBEFORE", (0, 0), (0, -1), 3, colors.HexColor(borde)),
                ("LEFTPADDING", (0, 0), (-1, -1), 9),
                ("RIGHTPADDING", (0, 0), (-1, -1), 9),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    return t


def _celda_valor(h: dict) -> list:
    celda = [_p(h["valor_texto"], "celda_num")]
    if h.get("personas_texto"):
        celda.append(_p(h["personas_texto"], "nota_der"))
    return celda


def pdf(t: dict) -> bytes:
    n = t["narrativa"]
    hechos = {h["id"]: h for h in n["hechos"]}
    ancho = LETTER[0] - 4 * cm
    historia = []

    # Portada
    historia += [
        _p(f"{n['area'].upper()} · {nombre_mes(n['periodo']).upper()}", "alcance"),
        _p("Reporte ejecutivo de Recursos Humanos", "titulo"),
        _p(linea_origen(n), "nota"),
        Spacer(1, 8),
    ]
    aprobacion = [_p(linea_aprobacion(t) + ".", "h3")]
    if t.get("comentario_revision"):
        aprobacion.append(_p(f"Comentario: {t['comentario_revision']}", "texto"))
    aprobacion.append(_p(f"Solicitado por {t.get('solicitada_por') or '—'}. Revisado por una persona de RRHH "
                         "distinta de quien lo solicitó.", "nota"))
    historia.append(_recuadro(aprobacion, "#e6f4e6", APROBADO, ancho))

    # Resumen ejecutivo
    historia += [_p("Resumen ejecutivo", "h2"), _p(n["resumen"], "resumen"), Spacer(1, 7)]
    fichas = [[_p(str(cuantos), "cifra"), Insignia(e, 8.5)] for e, cuantos in conteos(n)]
    resumen = Table([fichas], colWidths=[ancho / len(fichas)] * len(fichas))
    resumen.setStyle(TableStyle([
        *[("BACKGROUND", (i, 0), (i, 0), colors.HexColor(ESTADOS[e][2])) for i, (e, _) in enumerate(conteos(n))],
        ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor(BORDE)),
        ("INNERGRID", (0, 0), (-1, -1), 2, colors.white),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
    ]))
    historia.append(resumen)
    if n.get("mes_estable"):
        historia += [Spacer(1, 4), _p("Mes estable: sin alertas ni indicadores por vigilar.", "nota")]

    # Panel de indicadores clave (incluye la tendencia contra mes y año anterior)
    filas = [[_p(c, "encabezado") for c in ("Indicador", "Valor", "Vs mes anterior", "Vs año anterior", "Semáforo")]]
    for h in n["hechos"]:
        filas.append([
            _p(h["nombre"], "celda"),
            _celda_valor(h),
            _p(variacion(h.get("var_mes_ant_texto"), h.get("evolucion_mes_ant")), "celda"),
            _p(variacion(h.get("var_anio_ant_texto"), h.get("evolucion_anio_ant")), "celda"),
            Insignia(estado_visible(h)),
        ])
    historia += [
        _p("Indicadores clave", "h2"),
        _tabla(filas, [ancho * x for x in (0.27, 0.22, 0.19, 0.19, 0.13)]),
    ]
    if n.get("indicadores_sin_evaluar"):
        historia += [Spacer(1, 4), _p("Sin evaluar por muestra insuficiente (menos de 5 personas): "
                                      + ", ".join(n["indicadores_sin_evaluar"]) + ".", "nota")]

    # Hallazgos
    if n["hallazgos"]:
        historia.append(_p("Hallazgos", "h2"))
        for i, h in enumerate(n["hallazgos"], 1):
            historia.append(KeepTogether([
                _p(f"{i}. {h['titulo']}", "h3"),
                _p(h["texto"]),
                _p(f"{respaldo(h['hechos'], hechos, h['fuentes'])} {CONFIANZA.get(h['confianza'], h['confianza'])}.",
                   "nota"),
                Spacer(1, 6),
            ]))

    # Alertas y riesgos
    if alertas(n):
        filas = []
        for h in alertas(n):
            detalle = h["motivo_vigilancia"] if h.get("motivo_vigilancia") else f"el valor {h['situacion']}"
            filas.append([Insignia(estado_visible(h)), Paragraph(
                f"<b>{escape(h['nombre'])}</b>: {escape(h['valor_texto'])}; {escape(detalle)}.", E["celda"])])
        historia += [_p("Alertas y riesgos", "h2"),
                     _tabla(filas, [ancho * 0.16, ancho * 0.84], encabezado=False)]

    # Recomendaciones
    if n["recomendaciones"]:
        historia.append(_p("Recomendaciones", "h2"))
        for i, r in enumerate(n["recomendaciones"], 1):
            historia.append(KeepTogether([
                _p(f"{i}. {r['accion']}", "h3"),
                _p(f"{respaldo(r['hechos'], hechos, r['fuentes'])} {CONFIANZA.get(r['confianza'], r['confianza'])}.",
                   "nota"),
                Spacer(1, 6),
            ]))

    # Anexo de metodologia (pagina aparte: lo ejecutivo queda al frente)
    historia += [
        PageBreak(),
        _p("Anexo: metodología", "h2"),
        _p(f"Fecha de corte: datos al cierre de {nombre_mes(n['periodo'])}. Las cifras las calcula el motor "
           "analítico con reglas fijas, sin IA. La IA solo redacta el resumen, los hallazgos y las recomendaciones "
           "a partir de esas cifras; un validador automático rechaza cualquier cifra que no esté en los "
           "indicadores, y una persona de RRHH aprueba el texto antes de distribuirlo."),
        Spacer(1, 4),
        _p("Semáforo: Crítico si el valor cruzó el umbral crítico; Atención si cruzó el de atención; En rango si "
           "no cruzó ninguno; Por vigilar si está en rango pero empeoró de forma relevante. La confianza es alta, "
           "media o baja según el tamaño de la muestra y los meses de historia. Por privacidad, los indicadores de "
           "equipos de menos de 5 personas no se muestran."),
        Spacer(1, 8),
    ]
    filas = [[_p(c, "encabezado") for c in ("Indicador", "Fórmula", "Umbral atención / crítico",
                                            "Confianza", "Fuente")]]
    for h in n["hechos"]:
        filas.append([
            _p(h["nombre"], "celda"),
            _p(FORMULAS.get(h["indicador"], "—"), "nota"),
            _p(f"{h['umbral_atencion_texto']} / {h['umbral_critico_texto']}", "celda"),
            _p(f"{h['confianza']}: {h['motivo_confianza']}", "nota"),
            _p(h["fuente"], "nota"),
        ])
    historia.append(_tabla(filas, [ancho * x for x in (0.17, 0.25, 0.14, 0.17, 0.27)]))
    historia += [
        Spacer(1, 10),
        _p(f"Bitácora: reporte #{t['id']} · solicitado por {t.get('solicitada_por') or '—'} el "
           f"{fecha_hora(t['solicitada_en'])} · {linea_aprobacion(t).lower()} · "
           f"{n['intentos']} {'intento' if n['intentos'] == 1 else 'intentos'} del modelo.", "nota"),
    ]

    salida = io.BytesIO()
    doc = SimpleDocTemplate(
        salida, pagesize=LETTER, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.8 * cm, bottomMargin=2 * cm,
        title=f"Reporte ejecutivo de RRHH: {n['area']}, {nombre_mes(n['periodo'])}",
        author="Motor Inteligente de Reportes de RRHH", subject=linea_aprobacion(t), creator="rrhh-motor-reportes",
    )
    pie_der = f"Reporte #{t['id']} · {n['area']} · {nombre_mes(n['periodo'])}"
    doc.build(historia, canvasmaker=lambda *a, **k: CanvasNumerado(*a, pie_izq=PIE, pie_der=pie_der, **k))
    return salida.getvalue()


# -------------------------------------------------------------- presentacion
# 16:9. python-pptx no mide el texto: los tamanos de letra y la cantidad de
# elementos por diapositiva estan elegidos para que quepa, y los hallazgos (de
# largo variable) se reparten segun un alto estimado. Verificado en PowerPoint.
ANCHO, ALTO = Inches(13.333), Inches(7.5)
MARGEN = Inches(0.6)
ALTO_PARA_HALLAZGOS = 5.4  # pulgadas entre el titulo y el pie
ALERTAS_POR_DIAPOSITIVA = 7
INDICADORES_POR_DIAPOSITIVA = 10
COMENTARIO_EN_PORTADA = 180  # caracteres (dos lineas)


def _rgb(hexa: str) -> RGBColor:
    return RGBColor.from_string(hexa.lstrip("#"))


def _caja(diapositiva, x, y, ancho, alto, parrafos, ancla=MSO_ANCHOR.TOP):
    """parrafos: lista de (texto, tamano, negritas, color) o listas de esos (varias corridas en un parrafo)."""
    caja = diapositiva.shapes.add_textbox(x, y, ancho, alto)
    marco = caja.text_frame
    marco.word_wrap = True
    marco.vertical_anchor = ancla
    marco.margin_left = marco.margin_right = Inches(0.05)
    for i, parrafo in enumerate(parrafos):
        p = marco.paragraphs[0] if i == 0 else marco.add_paragraph()
        p.space_after = Pt(6)
        for texto, tamano, negritas, color in parrafo if isinstance(parrafo, list) else [parrafo]:
            r = p.add_run()
            r.text = texto
            r.font.size = Pt(tamano)
            r.font.bold = negritas
            r.font.color.rgb = _rgb(color)
    return caja


def _rectangulo(diapositiva, x, y, ancho, alto, relleno, forma=MSO_SHAPE.RECTANGLE):
    s = diapositiva.shapes.add_shape(forma, x, y, ancho, alto)
    s.fill.solid()
    s.fill.fore_color.rgb = _rgb(relleno)
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def _diapositiva(prs, titulo: str, pie: str):
    d = prs.slides.add_slide(prs.slide_layouts[6])  # en blanco
    _rectangulo(d, 0, 0, ANCHO, Inches(0.12), ACENTO)
    _caja(d, MARGEN, Inches(0.35), ANCHO - 2 * MARGEN, Inches(0.8), [(titulo, 28, True, TINTA)])
    _caja(d, MARGEN, ALTO - Inches(0.5), Inches(6), Inches(0.3), [(PIE, 10, False, SECUNDARIO)])
    _caja(d, ANCHO - MARGEN - Inches(6), ALTO - Inches(0.5), Inches(6), Inches(0.3), [(pie, 10, False, SECUNDARIO)])
    d.shapes[-1].text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    return d


def _insignia(d, x, y, estado, tamano=12):
    etiqueta, color, _ = ESTADOS.get(estado, (estado, SECUNDARIO, "#ffffff"))
    lado = Pt(tamano * 0.8)
    _rectangulo(d, x, y + Pt(tamano * 0.35), lado, lado, color, MSO_SHAPE.OVAL)
    _caja(d, x + lado + Inches(0.05), y, Inches(1.8), Pt(tamano * 1.8), [(etiqueta, tamano, True, TINTA)])


def _lineas(texto: str, por_linea: int) -> int:
    return max(1, math.ceil(len(texto) / por_linea))


def _alto_hallazgo(h: dict, hechos: dict) -> float:
    """Alto estimado en pulgadas. Medido en PowerPoint con 12 pulgadas de ancho:
    caben unos 118 caracteres por linea a 15 pt y 165 a 11 pt (se usa menos, por margen)."""
    lineas_respaldo = _lineas(respaldo(h["hechos"], hechos, h["fuentes"]), 150)
    return 0.45 + _lineas(h["texto"], 105) * 0.27 + lineas_respaldo * 0.2 + 0.35


def _agrupar(elementos: list, alto, disponible: float) -> list[list]:
    """Reparte en diapositivas segun el alto de cada elemento (la IA escribe textos de largo variable)."""
    grupos, actual, usado = [], [], 0.0
    for e in elementos:
        if actual and usado + alto(e) > disponible:
            grupos.append(actual)
            actual, usado = [], 0.0
        actual.append(e)
        usado += alto(e)
    return grupos + [actual] if actual else grupos


def _partes(lista, tamano):
    return [lista[i:i + tamano] for i in range(0, len(lista), tamano)] or [[]]


def _titulo_parte(titulo, i, total):
    return titulo if total == 1 else f"{titulo} ({i} de {total})"


def presentacion(t: dict) -> bytes:
    n = t["narrativa"]
    hechos = {h["id"]: h for h in n["hechos"]}
    pie = f"Reporte #{t['id']} · {n['area']} · {nombre_mes(n['periodo'])}"
    util = ANCHO - 2 * MARGEN
    prs = Presentation()
    prs.slide_width, prs.slide_height = ANCHO, ALTO
    prs.core_properties.title = f"Reporte ejecutivo de RRHH: {n['area']}, {nombre_mes(n['periodo'])}"
    prs.core_properties.author = "Motor Inteligente de Reportes de RRHH"
    prs.core_properties.subject = linea_aprobacion(t)

    # 1. Portada
    d = prs.slides.add_slide(prs.slide_layouts[6])
    _rectangulo(d, 0, 0, Inches(0.35), ALTO, ACENTO)
    _caja(d, Inches(1.1), Inches(2.0), Inches(11), Inches(0.5),
          [(f"{n['area'].upper()} · {nombre_mes(n['periodo']).upper()}", 16, True, ACENTO)])
    _caja(d, Inches(1.1), Inches(2.6), Inches(11), Inches(1.4),
          [("Reporte ejecutivo de Recursos Humanos", 40, True, TINTA)])
    _rectangulo(d, Inches(1.1), Inches(4.35), Inches(0.08), Inches(0.9), APROBADO)
    aprobacion = [(linea_aprobacion(t), 18, True, TINTA)]
    if t.get("comentario_revision"):
        comentario = t["comentario_revision"]
        if len(comentario) > COMENTARIO_EN_PORTADA:  # hasta 1000 caracteres: completo en el PDF
            comentario = comentario[:COMENTARIO_EN_PORTADA].rstrip() + "… (completo en el PDF)"
        aprobacion.append((f"Comentario: {comentario}", 13, False, SECUNDARIO))
    _caja(d, Inches(1.35), Inches(4.3), Inches(10.5), Inches(1.0), aprobacion)
    _caja(d, Inches(1.1), Inches(5.6), Inches(11), Inches(0.9), [(linea_origen(n), 12, False, SECUNDARIO)])
    _caja(d, Inches(1.1), ALTO - Inches(0.5), Inches(6), Inches(0.3), [(PIE, 10, False, SECUNDARIO)])

    # 2. Resumen ejecutivo
    d = _diapositiva(prs, "Resumen ejecutivo", pie)
    _caja(d, MARGEN, Inches(1.35), util, Inches(3.0), [(n["resumen"], 20, False, TINTA)])
    fichas = conteos(n)
    separacion = Inches(0.25)
    ancho_ficha = int((util - separacion * (len(fichas) - 1)) / len(fichas))
    for i, (estado, cuantos) in enumerate(fichas):
        x = MARGEN + i * (ancho_ficha + separacion)
        etiqueta, color, fondo = ESTADOS[estado]
        _rectangulo(d, x, Inches(4.6), ancho_ficha, Inches(1.9), fondo)
        _rectangulo(d, x, Inches(4.6), Inches(0.08), Inches(1.9), color)
        _caja(d, x + Inches(0.3), Inches(4.75), ancho_ficha - Inches(0.4), Inches(1.0), [(str(cuantos), 44, True, TINTA)])
        _insignia(d, x + Inches(0.3), Inches(5.8), estado, 14)
    if n.get("mes_estable"):
        _caja(d, MARGEN, Inches(6.55), util, Inches(0.35),
              [("Mes estable: sin alertas ni indicadores por vigilar.", 12, False, SECUNDARIO)])

    # 3. Indicadores clave
    partes = _partes(n["hechos"], INDICADORES_POR_DIAPOSITIVA)
    for i, parte in enumerate(partes, 1):
        d = _diapositiva(prs, _titulo_parte("Indicadores clave", i, len(partes)), pie)
        columnas = [("Indicador", 0.30), ("Valor", 0.15), ("Vs mes anterior", 0.2), ("Vs año anterior", 0.2),
                    ("Semáforo", 0.15)]
        alto_fila = Inches(0.46)
        tabla = d.shapes.add_table(len(parte) + 1, len(columnas), MARGEN, Inches(1.3), util,
                                   alto_fila * (len(parte) + 1)).table
        for j, (nombre, fraccion) in enumerate(columnas):
            tabla.columns[j].width = int(util * fraccion)
        for fila, h in enumerate([None, *parte]):
            valores = [c for c, _ in columnas] if h is None else [
                h["nombre"], h["valor_texto"],
                variacion(h.get("var_mes_ant_texto"), h.get("evolucion_mes_ant")),
                variacion(h.get("var_anio_ant_texto"), h.get("evolucion_anio_ant")),
                ESTADOS[estado_visible(h)][0],
            ]
            for j, valor in enumerate(valores):
                celda = tabla.cell(fila, j)
                celda.fill.solid()
                if h is None:
                    celda.fill.fore_color.rgb = _rgb(ACENTO)
                elif j == 4:
                    celda.fill.fore_color.rgb = _rgb(ESTADOS[estado_visible(h)][2])
                else:
                    celda.fill.fore_color.rgb = _rgb("#ffffff" if fila % 2 else "#f4f6f9")
                celda.vertical_anchor = MSO_ANCHOR.MIDDLE
                p = celda.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.RIGHT if j == 1 and h is not None else PP_ALIGN.LEFT
                r = p.add_run()
                r.text = valor
                r.font.size = Pt(13 if h is None else 14)
                r.font.bold = h is None or j == 4
                r.font.color.rgb = _rgb("#ffffff" if h is None else TINTA)
        if i == len(partes) and n.get("indicadores_sin_evaluar"):
            _caja(d, MARGEN, ALTO - Inches(0.95), util, Inches(0.35),
                  [("Sin evaluar por muestra insuficiente: " + ", ".join(n["indicadores_sin_evaluar"]) + ".",
                    12, False, SECUNDARIO)])

    # 4. Hallazgos
    def alto(h):
        return _alto_hallazgo(h, hechos)

    partes = _agrupar(n["hallazgos"], alto, ALTO_PARA_HALLAZGOS)
    numero = 0
    for i, parte in enumerate(partes, 1):
        d = _diapositiva(prs, _titulo_parte("Hallazgos", i, len(partes)), pie)
        y = 1.35
        for h in parte:
            numero += 1
            _rectangulo(d, MARGEN, Inches(y + 0.08), Inches(0.08), Inches(alto(h) - 0.4), ACENTO)
            _caja(d, MARGEN + Inches(0.25), Inches(y), util - Inches(0.25), Inches(alto(h) - 0.2), [
                [(f"{numero}. {h['titulo']}", 18, True, TINTA),
                 (f"   {CONFIANZA.get(h['confianza'], h['confianza'])}", 12, False, SECUNDARIO)],
                (h["texto"], 15, False, TINTA),
                (respaldo(h["hechos"], hechos, h["fuentes"]), 11, False, SECUNDARIO),
            ])
            y += alto(h)

    # 5. Alertas y riesgos
    lista = alertas(n)
    partes = _partes(lista, ALERTAS_POR_DIAPOSITIVA)
    for i, parte in enumerate(partes, 1):
        if not parte:
            break
        d = _diapositiva(prs, _titulo_parte("Alertas y riesgos", i, len(partes)), pie)
        for k, h in enumerate(parte):
            y = Inches(1.4) + k * Inches(0.75)
            _insignia(d, MARGEN, y, estado_visible(h), 14)
            detalle = h["motivo_vigilancia"] if h.get("motivo_vigilancia") else f"el valor {h['situacion']}"
            _caja(d, MARGEN + Inches(1.7), y, util - Inches(1.7), Inches(0.7),
                  [[(f"{h['nombre']}: {h['valor_texto']}", 16, True, TINTA), (f"; {detalle}.", 16, False, TINTA)]])

    # 6. Recomendaciones (de 1 a 3)
    if n["recomendaciones"]:
        d = _diapositiva(prs, "Recomendaciones", pie)
        for k, r in enumerate(n["recomendaciones"][:3]):
            y = Inches(1.35) + k * Inches(1.75)
            _rectangulo(d, MARGEN, y, Inches(0.6), Inches(0.6), ACENTO, MSO_SHAPE.OVAL)
            _caja(d, MARGEN, y, Inches(0.6), Inches(0.6), [(str(k + 1), 20, True, "#ffffff")], MSO_ANCHOR.MIDDLE)
            d.shapes[-1].text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            _caja(d, MARGEN + Inches(0.85), y - Inches(0.05), util - Inches(0.85), Inches(1.6), [
                (r["accion"], 18, True, TINTA),
                (f"{respaldo(r['hechos'], hechos, r['fuentes'])} {CONFIANZA.get(r['confianza'], r['confianza'])}.",
                 12, False, SECUNDARIO),
            ])

    # 7. Metodologia y aprobacion
    d = _diapositiva(prs, "Metodología y aprobación", pie)
    _caja(d, MARGEN, Inches(1.35), util, Inches(5.2), [
        (f"Fecha de corte: datos al cierre de {nombre_mes(n['periodo'])}.", 15, True, TINTA),
        ("Las cifras las calcula el motor analítico con reglas fijas, sin IA. La IA solo redacta el texto a partir "
         "de esas cifras; un validador automático rechaza cualquier cifra que no esté en los indicadores.",
         15, False, TINTA),
        ("Semáforo: Crítico si el valor cruzó el umbral crítico; Atención si cruzó el de atención; En rango si no "
         "cruzó ninguno; Por vigilar si está en rango pero empeoró de forma relevante. Por privacidad, los "
         "indicadores de equipos de menos de 5 personas no se muestran.", 15, False, TINTA),
        (linea_origen(n), 13, False, SECUNDARIO),
        (f"Solicitado por {t.get('solicitada_por') or '—'} el {fecha_hora(t['solicitada_en'])}. "
         f"{linea_aprobacion(t)}, una persona de RRHH distinta de quien lo solicitó. "
         "Fórmulas, umbrales y fuentes de cada indicador: anexo del reporte en PDF.", 13, False, SECUNDARIO),
    ])

    salida = io.BytesIO()
    prs.save(salida)
    return salida.getvalue()
