"""Pruebas de la exportacion a PDF y presentacion (paso 9, RF-08).

Dos partes: el armado de los archivos (sin base de datos) y las reglas de
la API: solo se exportan reportes aprobados, cada rol solo lo suyo y cada
descarga queda en la bitacora. Como en test_revision.py, se usan usuarios
reales (prefijo zz_exp_) y las narrativas se insertan directo en la tabla.
"""
import io
import json

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Emu
from pypdf import PdfReader
from sqlalchemy import text

from app import exportar, seguridad, trabajos
from app.database import engine
from app.main import app
from app.seguridad import crear_usuario, usuario_actual

CLAVE = "Clave-de-prueba-123"
PREFIJO = "zz_exp_"
LARGO = (
    "La rotación voluntaria subió a 3.4 % y supera lo observado en los últimos doce meses del área, con una "
    "diferencia de +1.9 pts frente al mes anterior y de +2.2 pts frente al mismo mes del año anterior. "
)


def hecho(id_, nombre, estado, valor_texto, **extra):
    return {
        "id": id_, "indicador": extra.pop("indicador", "rotacion_voluntaria"), "nombre": nombre, "unidad": "%",
        "area": "Ventas", "periodo": "2026-08", "valor": 3.4, "valor_texto": valor_texto,
        "var_mes_ant": 1.9, "var_mes_ant_texto": "+1.9 pts", "var_anio_ant": None, "var_anio_ant_texto": None,
        "estado": estado, "n": 120, "personas_texto": None, "umbral_atencion": 1.5, "umbral_atencion_texto": "1.5 %",
        "umbral_critico": 3.0, "umbral_critico_texto": "3.0 %", "sentido": "valores más altos son peores",
        "evolucion_mes_ant": "empeoró", "evolucion_anio_ant": None, "en_vigilancia": False, "motivo_vigilancia": None,
        "situacion": "cruzó el umbral crítico", "accion_base": None, "confianza": "alta",
        "motivo_confianza": "muestra e historia suficientes", "fuente": "HRIS (empleados)", **extra,
    }


def reporte(**cambios) -> dict:
    """Reporte aprobado en el peor caso: textos largos, todos los estados y caracteres especiales."""
    hechos = [
        hecho("H01", "Rotación voluntaria mensual", "rojo", "3.4 %"),
        hecho("H02", "eNPS", "amarillo", "8.0 puntos", indicador="enps", situacion="cruzó el umbral de atención"),
        hecho("H03", "Cobertura de capacitación", "verde", "61.0 %", indicador="cobertura_capacitacion",
              en_vigilancia=True, motivo_vigilancia="está cerca del umbral de atención",
              personas_texto="61 de 100 empleados activos se capacitaron (39 no)"),
        hecho("H04", "Tiempo de contratación", "verde", "30.0 días", indicador="tiempo_contratacion",
              confianza="baja", motivo_confianza="muestra muy pequeña (n=3)"),
    ]
    narrativa = {
        "periodo": "2026-08", "area_id": 1, "area": "Ventas", "proveedor": "ollama", "modelo": "qwen3:8b",
        "version_prompt": "v5", "generado_en": "2026-10-03T18:00:00Z", "requiere_revision": True,
        "mes_estable": False, "indicadores_sin_evaluar": ["Tasa de finalización de cursos"], "intentos": 2,
        "resumen": "En agosto de 2026 hay 1 indicador en rojo y 1 en amarillo: rotación voluntaria (3.4 %) y "
                   "eNPS (8.0 puntos). Conviene vigilar la cobertura de capacitación <61.0 %> & seguir el mes.",
        "hallazgos": [
            {"titulo": f"Hallazgo largo {i}", "texto": LARGO * 3, "hechos": ["H01", "H02"], "confianza": "alta",
             "fuentes": ["HRIS (empleados)", "Encuesta de clima (respuestas_clima)"]}
            for i in range(1, 6)
        ],
        "recomendaciones": [
            {"accion": "Explorar con el área las razones que las personas dan al renunciar y dar seguimiento "
                       "mensual con la gerencia, revisando también el clima del equipo en la siguiente encuesta.",
             "hechos": ["H01"], "confianza": "alta", "fuentes": ["HRIS (empleados)"]}
        ] * 3,
        "conteo_estados": {"rojo": 1, "amarillo": 1, "verde": 2, "por_vigilar": 1, "total": 4},
        "hechos": hechos, "advertencias": [],
    }
    narrativa.update(cambios)
    return {
        "id": 42, "area_id": 1, "periodo": "2026-08", "estado": "lista", "solicitada_en": "2026-10-03T17:55:00Z",
        "terminada_en": "2026-10-03T18:00:00Z", "solicitada_por": "ana", "revision": "aprobada",
        "revisada_por": "eva", "revisada_en": "2026-10-04T16:30:00Z",
        "comentario_revision": "Cifras revisadas contra el tablero <ok> & listas", "narrativa": narrativa,
    }


def texto_pdf(contenido: bytes) -> str:
    """Texto del PDF con los espacios normalizados (las celdas parten las lineas)."""
    return " ".join(" ".join(p.extract_text() for p in PdfReader(io.BytesIO(contenido)).pages).split())


def textos_pptx(contenido: bytes) -> list[str]:
    """Texto de cada diapositiva (cajas y tablas)."""
    salida = []
    for d in Presentation(io.BytesIO(contenido)).slides:
        partes = []
        for s in d.shapes:
            if s.has_text_frame and s.text_frame.text:
                partes.append(s.text_frame.text)
            if s.has_table:
                partes += [c.text for fila in s.table.rows for c in fila.cells]
        salida.append("\n".join(partes))
    return salida


# --------------------------------------------------------- armado de archivos
def test_pdf_tiene_las_secciones_del_reporte_y_la_aprobacion():
    contenido = exportar.generar(reporte(), "pdf")
    assert contenido.startswith(b"%PDF")
    t = texto_pdf(contenido)
    for seccion in ("Resumen ejecutivo", "Indicadores clave", "Hallazgos", "Alertas y riesgos", "Recomendaciones",
                    "Anexo: metodología"):
        assert seccion in t
    assert "VENTAS · AGOSTO 2026" in t
    assert "Aprobado por eva el 4 oct 2026, 10:30" in t  # hora de la Ciudad de Mexico (UTC-6)
    assert "Solicitado por ana" in t
    assert "Página 1 de" in t and "Confidencial" in t
    assert "61 de 100 empleados activos" in t  # conteo de personas en capacitacion
    assert "Tasa de finalización de cursos" in t  # sin evaluar
    assert "Por vigilar" in t and "Crítico" in t and "Atención" in t


def test_pdf_escapa_el_texto_de_la_ia_y_del_revisor():
    """'<' y '&' son marcado para ReportLab: sin escapar, el PDF fallaria o perderia texto."""
    t = texto_pdf(exportar.generar(reporte(), "pdf"))
    assert "<61.0 %> & seguir" in t
    assert "<ok> & listas" in t


def test_pdf_metadatos():
    info = PdfReader(io.BytesIO(exportar.generar(reporte(), "pdf"))).metadata
    assert info.title == "Reporte ejecutivo de RRHH: Ventas, agosto 2026"
    assert "Aprobado por eva" in info.subject


def test_pdf_sin_alertas_ni_recomendaciones():
    """Mes estable: no hay secciones vacias."""
    r = reporte(mes_estable=True, hallazgos=[], recomendaciones=[], indicadores_sin_evaluar=[],
                conteo_estados={"rojo": 0, "amarillo": 0, "verde": 1, "por_vigilar": 0, "total": 1})
    r["narrativa"]["hechos"] = [hecho("H01", "eNPS", "verde", "25.0 puntos", situacion="sin alerta")]
    r["comentario_revision"] = None
    t = texto_pdf(exportar.generar(r, "pdf"))
    assert "Mes estable" in t
    assert "Alertas y riesgos" not in t and "Recomendaciones" not in t and "Comentario:" not in t


def test_presentacion_tiene_las_diapositivas_del_reporte():
    contenido = exportar.generar(reporte(), "pptx")
    diapositivas = textos_pptx(contenido)
    todo = "\n".join(diapositivas)
    assert "Reporte ejecutivo de Recursos Humanos" in diapositivas[0]
    assert "Aprobado por eva el 4 oct 2026, 10:30" in diapositivas[0]
    for titulo in ("Resumen ejecutivo", "Indicadores clave", "Hallazgos", "Alertas y riesgos", "Recomendaciones",
                   "Metodología y aprobación"):
        assert titulo in todo
    assert "<61.0 %> & seguir" in todo  # en pptx el texto no es marcado: va tal cual
    assert Presentation(io.BytesIO(contenido)).core_properties.title == "Reporte ejecutivo de RRHH: Ventas, agosto 2026"


def test_presentacion_reparte_los_hallazgos_largos_y_nada_sale_de_la_diapositiva():
    contenido = exportar.generar(reporte(), "pptx")
    prs = Presentation(io.BytesIO(contenido))
    titulos = [t.split("\n")[0] for t in textos_pptx(contenido)]
    hallazgos = [t for t in titulos if t.startswith("Hallazgos")]
    assert len(hallazgos) >= 2 and hallazgos[0] == f"Hallazgos (1 de {len(hallazgos)})"
    todo = "\n".join(textos_pptx(contenido))
    assert all(f"{i}. Hallazgo largo {i}" in todo for i in range(1, 6))  # ninguno se pierde
    for d in prs.slides:
        for s in d.shapes:
            assert 0 <= s.left and s.left + s.width <= prs.slide_width + Emu(1)
            assert 0 <= s.top and s.top + s.height <= prs.slide_height + Emu(1)
    # Ningun hallazgo invade al siguiente
    for d in prs.slides:
        cajas = sorted((s for s in d.shapes if s.has_text_frame and "Hallazgo largo" in s.text_frame.text),
                       key=lambda s: s.top)
        for a, b in zip(cajas, cajas[1:]):
            assert a.top + a.height <= b.top


def test_nombre_de_archivo_solo_ascii():
    r = reporte()
    r["narrativa"]["area"] = "Logística y Almacén"
    assert exportar.nombre_archivo(r, "pdf") == "reporte-rrhh-logistica-y-almacen-2026-08.pdf"


def test_formato_desconocido():
    with pytest.raises(ValueError):
        exportar.generar(reporte(), "docx")


# ------------------------------------------------------------------- API
@pytest.fixture(autouse=True)
def sin_atajo_de_sesion(monkeypatch):
    app.dependency_overrides.pop(usuario_actual, None)
    monkeypatch.setenv("JWT_SECRETO", "s" * 48)


@pytest.fixture(scope="module")
def usuarios():
    seguridad.asegurar_tabla()
    with engine.begin() as conn:
        conn.execute(text("DELETE FROM usuarios WHERE usuario LIKE :p"), {"p": PREFIJO + "%"})
        ventas = conn.execute(text("SELECT id FROM areas WHERE nombre = 'Ventas'")).scalar_one()
    datos = {"rrhh": None, "rrhh2": None, "direccion": None, "gerente": ventas, "admin_ti": None}
    for nombre, area_id in datos.items():
        crear_usuario(PREFIJO + nombre, f"Prueba {nombre}", nombre.rstrip("2"), CLAVE, area_id)
    yield {nombre: PREFIJO + nombre for nombre in datos}
    with engine.begin() as conn:
        conn.execute(text("DELETE FROM usuarios WHERE usuario LIKE :p"), {"p": PREFIJO + "%"})


@pytest.fixture()
def client():
    trabajos.asegurar_tabla()
    with engine.connect() as conn:
        ultimo = conn.execute(text("SELECT COALESCE(MAX(id), 0) FROM narrativas")).scalar()
    yield TestClient(app)
    with engine.begin() as conn:  # las exportaciones se borran en cascada
        conn.execute(text("DELETE FROM narrativas WHERE id > :u"), {"u": ultimo})


@pytest.fixture()
def como(client, usuarios):
    def _como(nombre):
        r = client.post("/auth/login", data={"username": usuarios[nombre], "password": CLAVE})
        assert r.status_code == 200, r.text
        return {"Authorization": f"Bearer {r.json()['access_token']}"}
    return _como


@pytest.fixture()
def narrativa(client, usuarios):
    """Inserta una narrativa lista del consolidado con su resultado y devuelve su id."""
    def _crear(revision="aprobada", area_id=0):
        n = reporte(area_id=area_id, area="Corporativo")["narrativa"]
        with engine.begin() as conn:
            return conn.execute(
                text(
                    "INSERT INTO narrativas (area_id, periodo, estado, solicitada_por, terminada_en, resultado, "
                    "version_prompt, revision, revisada_por, revisada_en) "
                    "VALUES (:a, '2026-08-01', 'lista', :s, now(), CAST(:r AS jsonb), 'v5', :rev, "
                    "CASE WHEN :rev = 'pendiente' THEN NULL ELSE :revisor END, "
                    "CASE WHEN :rev = 'pendiente' THEN NULL ELSE now() END) RETURNING id"
                ),
                {"a": area_id, "s": usuarios["rrhh"], "r": json.dumps(n, ensure_ascii=False), "rev": revision,
                 "revisor": usuarios["rrhh2"]},
            ).scalar_one()
    return _crear


def exportaciones(id_):
    with engine.connect() as conn:
        return [tuple(f) for f in conn.execute(
            text("SELECT usuario, formato FROM exportaciones WHERE narrativa_id = :i ORDER BY id"), {"i": id_}
        ).all()]


@pytest.mark.parametrize("formato,inicio", [("pdf", b"%PDF"), ("pptx", b"PK")])
def test_exporta_un_reporte_aprobado_y_queda_en_la_bitacora(client, como, narrativa, usuarios, formato, inicio):
    id_ = narrativa()
    r = client.get(f"/narrativas/{id_}/exportar", params={"formato": formato}, headers=como("direccion"))
    assert r.status_code == 200, r.text
    assert r.content.startswith(inicio)
    assert r.headers["content-type"] == exportar.FORMATOS[formato]
    assert f'filename="reporte-rrhh-corporativo-2026-08.{formato}"' in r.headers["content-disposition"]
    assert exportaciones(id_) == [(usuarios["direccion"], formato)]


def test_el_pdf_de_la_api_dice_quien_aprobo(client, como, narrativa, usuarios):
    id_ = narrativa()
    r = client.get(f"/narrativas/{id_}/exportar", headers=como("rrhh"))  # pdf por defecto
    assert r.status_code == 200
    t = texto_pdf(r.content)
    assert f"Aprobado por {usuarios['rrhh2']}" in t and f"Reporte #{id_}" in t


@pytest.mark.parametrize("revision", ["pendiente", "rechazada"])
def test_solo_se_exportan_reportes_aprobados(client, como, narrativa, revision):
    id_ = narrativa(revision=revision)
    r = client.get(f"/narrativas/{id_}/exportar", headers=como("rrhh"))
    assert r.status_code == 409
    assert "aprobad" in r.json()["detail"]
    assert exportaciones(id_) == []


def test_no_se_exporta_lo_que_no_esta_terminado(client, como, usuarios):
    with engine.begin() as conn:
        id_ = conn.execute(
            text("INSERT INTO narrativas (area_id, periodo, estado, solicitada_por) "
                 "VALUES (0, '2026-08-01', 'en_proceso', :s) RETURNING id"), {"s": usuarios["rrhh"]}
        ).scalar_one()
    assert client.get(f"/narrativas/{id_}/exportar", headers=como("rrhh")).status_code == 409


def test_cada_rol_exporta_solo_lo_suyo(client, como, narrativa):
    consolidado = narrativa()
    assert client.get(f"/narrativas/{consolidado}/exportar", headers=como("admin_ti")).status_code == 403
    assert client.get(f"/narrativas/{consolidado}/exportar", headers=como("gerente")).status_code == 403
    assert exportaciones(consolidado) == []


def test_formato_invalido_e_inexistente(client, como, narrativa):
    id_ = narrativa()
    assert client.get(f"/narrativas/{id_}/exportar", params={"formato": "docx"}, headers=como("rrhh")).status_code == 422
    assert client.get("/narrativas/999999999/exportar", headers=como("rrhh")).status_code == 404


def test_sin_sesion_no_se_exporta(client, narrativa):
    assert client.get(f"/narrativas/{narrativa()}/exportar").status_code == 401


def test_comentario_largo_se_recorta_con_aviso_en_la_presentacion_y_completo_en_el_pdf():
    r = reporte()
    r["comentario_revision"] = "Revisado. " * 60 + "FIN"
    portada = textos_pptx(exportar.generar(r, "pptx"))[0]
    assert "… (completo en el PDF)" in portada and "FIN" not in portada
    pdf = texto_pdf(exportar.generar(r, "pdf"))
    assert pdf.count("Revisado.") == 60 and "FIN" in pdf
